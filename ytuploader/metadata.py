"""Bulk metadata editing: export your channel to a spreadsheet, edit, apply.

The dangerous part of videos.update is that it REPLACES the parts you send.
Omit `tags` from a snippet update and YouTube deletes the video's tags. So every
update here starts from the video's current snippet and overlays only the cells
you actually filled in.
"""

from __future__ import annotations

import csv
import difflib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from googleapiclient.errors import HttpError

from .core import norm_name, resolve_category

# Columns that identify which video a row refers to.
MATCH_COLS = ("video_id", "url", "file", "current_title")
# Columns that carry new values. Blank means "leave unchanged".
EDIT_COLS = ("title", "description", "tags", "category", "privacy", "publish_at")

HEADERS = ("video_id", "url", "file", "current_title",
           "title", "description", "tags", "category", "privacy", "publish_at")

# Written into the export so a human can see what they're changing from.
CURRENT_COLS = ("current_description", "current_tags", "current_category",
                "current_privacy", "duration", "published_at")


# Spreadsheets come from all over. Map common human headers onto our columns so
# a sheet doesn't have to be restructured before it can be used.
ALIASES = {
    "file": "file", "local filename": "file", "local filename (.mkv)": "file",
    "filename": "file", "file name": "file", "local file": "file", "path": "file",

    "video_id": "video_id", "video id": "video_id", "videoid": "video_id", "id": "video_id",

    "url": "url", "link": "url", "youtube url": "url", "watch url": "url",
    "new url": "url",

    "current_title": "current_title", "current title": "current_title",
    "existing title": "current_title", "old title": "current_title",

    "title": "title", "new title": "title", "youtube title": "title",
    "youtube title (to use)": "title", "title (to use)": "title",

    "description": "description", "new description": "description",
    "suggested description": "description", "description (to use)": "description",

    "tags": "tags", "new tags": "tags", "keywords": "tags",
    "category": "category", "new category": "category",
    "privacy": "privacy", "visibility": "privacy", "new privacy": "privacy",
    "publish_at": "publish_at", "publish at": "publish_at", "schedule": "publish_at",

    # Points at the *original* upload on the old channel. Never a match key — it
    # identifies a different video — but a good source of the original metadata.
    "original video url (reference)": "reference_url",
    "original url": "reference_url", "reference url": "reference_url",
    "original video url": "reference_url",
}

# Columns that describe the *old* video, not the one being edited. Treating one
# of these as a match key would send edits to the wrong channel entirely.
IGNORE_MARKERS = ("original", "reference", "confidence", "note", "match")


def canon_header(h: str) -> tuple[str, str]:
    """Return (canonical column, reason-if-ignored) for a spreadsheet header."""
    raw = " ".join(str(h or "").split()).strip()
    low = raw.lower()
    if not low:
        return "", ""
    if low in ALIASES and ALIASES[low] == "reference_url":
        return "reference_url", ""
    if any(m in low for m in IGNORE_MARKERS):
        return "", raw
    return ALIASES.get(low, low), ""


def _map_row(header: list[str], values) -> dict:
    out = {}
    for i, h in enumerate(header):
        if i >= len(values):
            break
        col, ignored = canon_header(h)
        if not col or ignored:
            continue
        v = values[i]
        out[col] = "" if v is None else str(v).strip()
    return out


def sheet_columns(path: Path) -> tuple[list[str], list[str]]:
    """(recognised canonical columns, headers that were ignored)."""
    header = _read_header(path)
    keep, dropped = [], []
    for h in header:
        col, ignored = canon_header(h)
        if ignored:
            dropped.append(ignored)
        elif col:
            keep.append(col)
    return keep, dropped


@dataclass
class Video:
    video_id: str
    title: str = ""
    description: str = ""
    tags: list[str] = field(default_factory=list)
    category: str = ""
    privacy: str = ""
    publish_at: str = ""
    duration: str = ""
    published_at: str = ""


@dataclass
class Edit:
    row: int
    video: Video | None = None
    changes: dict = field(default_factory=dict)   # field -> (old, new)
    problem: str = ""
    kind: str = ""        # "" | "pending" (not uploaded yet) | "problem"
    label: str = ""
    matched_by: str = ""  # video_id | url | file | title | fuzzy title | description
    confidence: float = 1.0


# --------------------------------------------------------------- read channel

def fetch_channel_videos(youtube, log=print) -> list[Video]:
    """Every video on the authenticated user's channel, with its current metadata."""
    ch = youtube.channels().list(part="contentDetails", mine=True).execute()
    items = ch.get("items") or []
    if not items:
        raise RuntimeError("No channel found for this account")
    uploads = items[0]["contentDetails"]["relatedPlaylists"]["uploads"]

    ids, page = [], None
    while True:
        res = youtube.playlistItems().list(
            part="contentDetails", playlistId=uploads,
            maxResults=50, pageToken=page).execute()
        ids += [i["contentDetails"]["videoId"] for i in res.get("items", [])]
        page = res.get("nextPageToken")
        if not page:
            break
    log(f"channel has {len(ids)} video(s)")

    out: list[Video] = []
    for i in range(0, len(ids), 50):
        res = youtube.videos().list(
            part="snippet,status,contentDetails",
            id=",".join(ids[i:i + 50]), maxResults=50).execute()
        for it in res.get("items", []):
            sn, st = it.get("snippet", {}), it.get("status", {})
            out.append(Video(
                video_id=it["id"],
                title=sn.get("title", ""),
                description=sn.get("description", ""),
                tags=sn.get("tags", []) or [],
                category=sn.get("categoryId", ""),
                privacy=st.get("privacyStatus", ""),
                publish_at=st.get("publishAt", "") or "",
                duration=it.get("contentDetails", {}).get("duration", ""),
                published_at=sn.get("publishedAt", ""),
            ))
    return out


# --------------------------------------------------------------- spreadsheets

def _ext(p: Path) -> str:
    return p.suffix.lower()


def export_sheet(path: Path, videos: list[Video], state_data: dict | None = None) -> None:
    """Write an editable sheet: identity + current values + blank columns to fill."""
    by_id = {}
    for key, rec in (state_data or {}).items():
        if isinstance(rec, dict) and rec.get("video_id"):
            by_id[rec["video_id"]] = key

    header = list(HEADERS) + list(CURRENT_COLS)
    rows = []
    for v in videos:
        rows.append([
            v.video_id, f"https://youtu.be/{v.video_id}", by_id.get(v.video_id, ""),
            v.title,
            "", "", "", "", "", "",                      # blank edit columns
            v.description, ", ".join(v.tags), v.category, v.privacy,
            v.duration, v.published_at,
        ])

    if _ext(path) in (".xlsx", ".xlsm"):
        _write_xlsx(path, header, rows)
    else:
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)


def _write_xlsx(path: Path, header: list[str], rows: list[list]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Excel support needs openpyxl — pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Videos"
    ws.append(header)

    edit_start = HEADERS.index("title") + 1
    edit_end = HEADERS.index("publish_at") + 1
    fill_edit = PatternFill("solid", fgColor="FFF3D6")
    fill_id = PatternFill("solid", fgColor="EDEDED")

    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
        if edit_start <= c <= edit_end:
            cell.fill = fill_edit
        elif c > edit_end:
            cell.fill = fill_id

    for r in rows:
        ws.append(r)

    widths = {"video_id": 14, "url": 26, "file": 40, "current_title": 46,
              "title": 46, "description": 46, "tags": 26, "category": 10,
              "privacy": 10, "publish_at": 20, "current_description": 46,
              "current_tags": 26, "current_category": 10, "current_privacy": 10,
              "duration": 11, "published_at": 22}
    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)

    ws.freeze_panes = "E2"
    wb.save(path)


def _read_header(path: Path) -> list[str]:
    if _ext(path) in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Excel support needs openpyxl — pip install openpyxl")
        ws = load_workbook(path, data_only=True, read_only=True).active
        for row in ws.iter_rows(values_only=True):
            return [str(h) if h is not None else "" for h in row]
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return next(csv.reader(f), [])


def read_sheet(path: Path) -> list[dict]:
    """Read a sheet into canonical column names, ignoring reference-only columns."""
    if _ext(path) in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Excel support needs openpyxl — pip install openpyxl")
        ws = load_workbook(path, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(it, [])]
        rows = [_map_row(header, r) for r in it]
    else:
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = next(reader, [])
            rows = [_map_row(header, r) for r in reader]

    return [r for r in rows if any(r.get(c) for c in MATCH_COLS)]


# --------------------------------------------------------------- matching

def _video_id_from(text: str) -> str:
    t = text.strip()
    if not t:
        return ""
    if "youtu" in t:
        for marker in ("v=", "youtu.be/", "/shorts/", "/embed/"):
            if marker in t:
                t = t.split(marker, 1)[1]
                break
        for sep in ("&", "?", "/"):
            t = t.split(sep, 1)[0]
    return t.strip()


def resolve(rows: list[dict], videos: list[Video], state_data: dict | None = None,
            match_by: str = "auto", fuzzy: float = 0.0) -> list[Edit]:
    """Work out which video each row means, and what would actually change.

    match_by: "auto" tries video_id → url → file → current_title. "title" also
    accepts the `title` column as the video's existing title (useful when a sheet
    lists the same title it already has). "description" matches on description
    text, for sheets that carry no usable title.

    fuzzy: 0 disables approximate title matching. Otherwise a similarity ratio
    (0.85–0.95 is sensible); the best match above it is used and reported with
    its score so you can check it in preview before applying.
    """
    by_id = {v.video_id: v for v in videos}
    by_title: dict[str, list[Video]] = {}
    by_desc: dict[str, list[Video]] = {}
    for v in videos:
        by_title.setdefault(norm_name(v.title), []).append(v)
        if v.description.strip():
            by_desc.setdefault(norm_name(v.description), []).append(v)

    file_to_id = {}
    for key, rec in (state_data or {}).items():
        if isinstance(rec, dict) and rec.get("video_id"):
            file_to_id[norm_name(Path(key).name)] = rec["video_id"]

    edits: list[Edit] = []
    for n, row in enumerate(rows, start=2):
        e = Edit(row=n)

        vid = _video_id_from(row.get("video_id", "")) or _video_id_from(row.get("url", ""))
        if vid and vid in by_id:
            e.video, e.matched_by = by_id[vid], "video_id"
        elif vid:
            e.problem = f"video id {vid} is not on this channel"

        tried_file = False
        if e.video is None and not e.problem and row.get("file"):
            tried_file = True
            got = file_to_id.get(norm_name(Path(row["file"]).name))
            if got and got in by_id:
                e.video, e.matched_by = by_id[got], "file"

        # Which text should stand in for "the video's current title"?
        title_key = row.get("current_title") or (
            row.get("title") if match_by in ("title", "description") else "")

        if e.video is None and not e.problem and title_key:
            hits = by_title.get(norm_name(title_key), [])
            if len(hits) == 1:
                e.video, e.matched_by = hits[0], "title"
            elif len(hits) > 1:
                e.problem = f"{len(hits)} videos share this title — use video_id instead"
            elif fuzzy:
                key = norm_name(title_key)
                best, score = None, 0.0
                for cand_key, cands in by_title.items():
                    if len(cands) != 1:
                        continue
                    r = difflib.SequenceMatcher(None, key, cand_key).ratio()
                    if r > score:
                        best, score = cands[0], r
                if best is not None and score >= fuzzy:
                    e.video, e.matched_by, e.confidence = best, "fuzzy title", score

        if (e.video is None and not e.problem and match_by == "description"
                and row.get("description")):
            hits = by_desc.get(norm_name(row["description"]), [])
            if len(hits) == 1:
                e.video, e.matched_by, e.problem = hits[0], "description", ""
            elif len(hits) > 1:
                e.problem = f"{len(hits)} videos share this description"

        if e.video is None and not e.problem:
            if tried_file:
                e.problem = "not uploaded yet (no entry in uploaded.json)"
                e.kind = "pending"
            elif title_key:
                e.problem = "no video on the channel has this title"
            else:
                e.problem = "row has no video_id, url, file or current_title"

        if e.problem and not e.kind:
            e.kind = "problem"

        if e.video is None:
            e.label = row.get("current_title") or row.get("file") or row.get("video_id") or "?"
            edits.append(e)
            continue

        e.label = e.video.title
        v = e.video
        for col in EDIT_COLS:
            new = (row.get(col) or "").strip()
            if not new:
                continue
            if col == "tags":
                old_list = v.tags
                new_list = [t.strip() for t in new.split(",") if t.strip()]
                if old_list != new_list:
                    e.changes["tags"] = (", ".join(old_list), ", ".join(new_list))
            elif col == "category":
                try:
                    new_id = resolve_category(new)
                except ValueError as ex:
                    e.problem = str(ex)
                    continue
                if new_id != v.category:
                    e.changes["category"] = (v.category, new_id)
            elif col == "privacy":
                if new.lower() not in ("private", "unlisted", "public"):
                    e.problem = f"privacy must be private, unlisted or public (got '{new}')"
                elif new.lower() != v.privacy:
                    e.changes["privacy"] = (v.privacy, new.lower())
            elif col == "publish_at":
                if new != v.publish_at:
                    e.changes["publish_at"] = (v.publish_at, new)
            elif col == "title":
                if len(new) > 100:
                    e.problem = f"title is {len(new)} characters (max 100)"
                elif new != v.title:
                    e.changes["title"] = (v.title, new)
            else:
                if new != getattr(v, col):
                    e.changes[col] = (getattr(v, col), new)

        edits.append(e)
    return edits


# --------------------------------------------------------------- apply

def apply_edit(youtube, e: Edit) -> None:
    """Send one update, preserving every field the sheet did not touch."""
    v = e.video
    assert v is not None

    snippet = {
        "title": e.changes.get("title", (None, v.title))[1],
        "description": e.changes.get("description", (None, v.description))[1],
        "tags": [t.strip() for t in e.changes["tags"][1].split(",") if t.strip()]
                if "tags" in e.changes else v.tags,
        "categoryId": e.changes.get("category", (None, v.category))[1] or "22",
    }
    body = {"id": v.video_id, "snippet": snippet}
    parts = ["snippet"]

    if "privacy" in e.changes or "publish_at" in e.changes:
        status = {"privacyStatus": e.changes.get("privacy", (None, v.privacy))[1]}
        pub = e.changes.get("publish_at", (None, v.publish_at))[1]
        if pub:
            status["privacyStatus"] = "private"
            status["publishAt"] = pub
        body["status"] = status
        parts.append("status")

    youtube.videos().update(part=",".join(parts), body=body).execute()

    # keep the in-memory copy honest so a second pass sees no diff
    for f, (_old, new) in e.changes.items():
        if f == "tags":
            v.tags = [t.strip() for t in new.split(",") if t.strip()]
        elif f == "category":
            v.category = new
        else:
            setattr(v, f, new)


def apply_all(youtube, edits: Iterable[Edit], log=print) -> dict:
    ok, failed, skipped = 0, [], 0
    for e in edits:
        if e.video is None or e.problem or not e.changes:
            skipped += 1
            continue
        try:
            apply_edit(youtube, e)
            ok += 1
            log(f"updated {e.video.video_id}  {', '.join(e.changes)}  — {e.label[:60]}")
        except HttpError as ex:
            failed.append((e.label, str(ex)[:200]))
            log(f"FAILED {e.video.video_id}: {ex}")
    return {"updated": ok, "failed": failed, "skipped": skipped}


# --------------------------------------------------------------- enrichment

def fetch_videos_by_id(youtube, ids: list[str], log=print) -> dict[str, Video]:
    """Look up any public videos by ID — including ones on another channel."""
    out: dict[str, Video] = {}
    ids = [i for i in dict.fromkeys(ids) if i]
    for i in range(0, len(ids), 50):
        chunk = ids[i:i + 50]
        res = youtube.videos().list(
            part="snippet,contentDetails", id=",".join(chunk), maxResults=50).execute()
        for it in res.get("items", []):
            sn = it.get("snippet", {})
            out[it["id"]] = Video(
                video_id=it["id"],
                title=sn.get("title", ""),
                description=sn.get("description", ""),
                tags=sn.get("tags", []) or [],
                category=sn.get("categoryId", ""),
                published_at=sn.get("publishedAt", ""),
                duration=it.get("contentDetails", {}).get("duration", ""),
            )
        log(f"  fetched {len(out)}/{len(ids)}")
    return out


# --------------------------------------------------------------- title format

_DATE8 = re.compile(r"^\s*(\d{8})\b")
_DATE6 = re.compile(r"^\s*(\d{6})\b")
_QUALITY = re.compile(r"[\[\(]\s*(\d+K(?:\s+\w+)?)\s*[\]\)]|(?<![\w\[\(])(\d+K)\b")


def extract_date(*sources: str) -> str:
    """Pull an event date (YYYYMMDD) out of a title or filename prefix."""
    for s in sources:
        s = s or ""
        m = _DATE8.match(s)
        if m:
            d = m.group(1)
            if "19000101" < d < "21000101":
                return d
        m = _DATE6.match(s)
        if m:                       # 230228 → 20230228
            return "20" + m.group(1)
    return ""


def normalize_title(title: str, date: str = "", keep_date: bool = True) -> str:
    """Tidy a title into: [YYYYMMDD ]Artist - Song @ Event [4K]"""
    t = " ".join(str(title or "").split())

    quality = ""
    def _grab(m):
        nonlocal quality
        quality = (m.group(1) or m.group(2) or "").upper()
        return " "
    t = _QUALITY.sub(_grab, t)

    t = _DATE8.sub("", t, count=1)
    t = _DATE6.sub("", t, count=1)

    t = re.sub(r"\s*@\s*", " @ ", t)
    t = re.sub(r"\s+-\s+|\s+–\s+", " - ", t)
    t = re.sub(r"\s*\|\s*", " | ", t)
    t = " ".join(t.split()).strip(" -|@")

    parts = []
    if keep_date and date:
        parts.append(date)
    parts.append(t)
    out = " ".join(p for p in parts if p)
    if quality:
        out = f"{out} [{quality}]"
    return " ".join(out.split())[:100]


def enrich_sheet(youtube, in_path: Path, out_path: Path, keep_date: bool = True,
                 use_publish_date: bool = False, only: str = "both", log=print) -> dict:
    """Rebuild a sheet's title/description from the original uploads it references.

    Reads the reference_url column, fetches those videos' real metadata, and
    writes normalised titles and the original descriptions into a new sheet.
    Nothing is sent to YouTube — this only produces a file for you to review.
    """
    rows = read_sheet(in_path)
    ids = [_video_id_from(r.get("reference_url", "")) for r in rows]
    found = fetch_videos_by_id(youtube, [i for i in ids if i], log)

    stats = {"rows": len(rows), "matched": 0, "no_reference": 0, "not_found": 0,
             "dated_from_name": 0, "dated_from_publish": 0, "undated": 0}
    out_rows = []

    for r, ref_id in zip(rows, ids):
        orig = found.get(ref_id) if ref_id else None
        if not ref_id:
            stats["no_reference"] += 1
        elif orig is None:
            stats["not_found"] += 1
        else:
            stats["matched"] += 1

        src_title = (orig.title if orig else "") or r.get("title", "")
        date = extract_date(r.get("file", ""), src_title, r.get("title", ""))
        if date:
            stats["dated_from_name"] += 1
        elif use_publish_date and orig and orig.published_at:
            date = orig.published_at[:10].replace("-", "")
            stats["dated_from_publish"] += 1
        elif keep_date:
            stats["undated"] += 1

        new_title = normalize_title(src_title, date, keep_date)
        new_desc = (orig.description if orig else "") or r.get("description", "")

        out_rows.append({
            "file": r.get("file", ""),
            "reference_url": r.get("reference_url", ""),
            # A blank cell means "leave this field alone" downstream.
            "title": new_title if only in ("both", "titles") else "",
            "description": new_desc if only in ("both", "descriptions") else "",
            "current_title": r.get("title", ""),
            "original_published": (orig.published_at[:10] if orig and orig.published_at else ""),
            "date_source": ("filename" if extract_date(r.get("file", ""), src_title) else
                            ("publish date" if date else "none")),
        })

    header = ["file", "reference_url", "title", "description",
              "current_title", "original_published", "date_source"]
    table = [[row[h] for h in header] for row in out_rows]
    if _ext(out_path) in (".xlsx", ".xlsm"):
        _write_xlsx(out_path, header, table)
    else:
        with out_path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(table)
    return stats
